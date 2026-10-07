"""
Document Export (Step 05)
Converts the fused JSON (layout + text) into a properly formatted Microsoft Word document.
"""

import json
import argparse
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

try:
    from src.utils.juktakkhor import restore_juktakkhor
except ModuleNotFoundError:
    from utils.juktakkhor import restore_juktakkhor

def set_bangla_font(run, font_name="Nirmala UI", size_pt=14, bold=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)

def export_to_docx(fused_json_path: str, output_docx_path: str):
    json_path = Path(fused_json_path).resolve()
    out_path = Path(output_docx_path).resolve()
    
    if not json_path.exists():
        print(f"Error: Fused JSON not found at {json_path}")
        return
        
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
        
    doc = Document()
    
    # Title
    doc_name = data.get("document_name", "Digitized Document")
    title = doc.add_heading(f"Digitized Document: {doc_name}", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if title.runs:
        set_bangla_font(title.runs[0], "Nirmala UI", 18, bold=True)
    
    # Add metadata
    meta = doc.add_paragraph()
    meta_run = meta.add_run(f"Total Regions Detected: {data.get('total_regions', 0)}\n")
    meta_run.italic = True
    set_bangla_font(meta_run, "Nirmala UI", 11)
    
    h2 = doc.add_heading("Recognized Text", level=2)
    if h2.runs:
        set_bangla_font(h2.runs[0], "Nirmala UI", 14, bold=True)
    
    # Sort regions by reading order
    regions = data.get("regions", [])
    regions.sort(key=lambda x: x.get("reading_order", 0))
    for region in regions:
        text = restore_juktakkhor(region.get("text", ""))
        if text.strip():
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            set_bangla_font(run, "Nirmala UI", 14)
    
    out_path.parent.mkdir(parents=True, exist_ok=True)
    saved_path = out_path
    try:
        doc.save(str(out_path))
        print(f"Successfully exported digitized document to: {out_path}")
    except PermissionError:
        fallback_path = out_path.with_name(f"{out_path.stem}_restored{out_path.suffix}")
        doc.save(str(fallback_path))
        print(f"\n[!] Notice: '{out_path.name}' is locked by Microsoft Word.")
        print(f"    Saved updated document as: {fallback_path}")
        saved_path = fallback_path
    return saved_path

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Export fused JSON to DOCX")
    parser.add_argument("--input", type=str, required=True, help="Path to fused JSON (Step 4)")
    parser.add_argument("--output", type=str, required=True, help="Path for output .docx file")
    args = parser.parse_args()
    
    export_to_docx(args.input, args.output)
